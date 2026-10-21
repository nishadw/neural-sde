"""Build a practice script PDF: each presented slide's thumbnail beside its speaker notes.

Usage: python practice_script.py <deck.pptx> <thumb_dir> <out.pdf> [n_presented]
Thumbnails: pdftoppm -png -r 90 deck.pdf <thumb_dir>/s
"""
import sys, re, html
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Image, Table, TableStyle, Spacer, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

D = "/usr/share/fonts/truetype/dejavu/"
pdfmetrics.registerFont(TTFont("DV", D + "DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("DVB", D + "DejaVuSans-Bold.ttf"))
pdfmetrics.registerFont(TTFont("DVI", D + "DejaVuSans-Oblique.ttf"))
INK, RED, MUTED = colors.HexColor("#1F2328"), colors.HexColor("#C8102E"), colors.HexColor("#6B7280")

deck, thumbs, out = sys.argv[1:4]
n = int(sys.argv[4]) if len(sys.argv) > 4 else 22
P = Presentation(deck)
st = dict(
    say=ParagraphStyle("say", fontName="DV", fontSize=9, leading=12, textColor=INK, leftIndent=9, firstLineIndent=-9, spaceAfter=3),
    ask=ParagraphStyle("ask", fontName="DVI", fontSize=8, leading=10.5, textColor=MUTED, spaceAfter=3),
    plan=ParagraphStyle("plan", fontName="DVB", fontSize=8.5, leading=11, textColor=RED, spaceAfter=3),
    head=ParagraphStyle("head", fontName="DVB", fontSize=10, leading=13, textColor=INK, spaceAfter=4),
    title=ParagraphStyle("title", fontName="DVB", fontSize=16, leading=20, textColor=INK, spaceAfter=4),
    sub=ParagraphStyle("sub", fontName="DV", fontSize=9, leading=12, textColor=MUTED, spaceAfter=10),
)
story = [Paragraph("Neural SDEs: Practice Script", st["title"]),
         Paragraph("25:00 slot = 20:00 talk + 5:00 discussion. Black = say it. Gray = only if asked. "
                   "Practice with a timer and check the timestamps.", st["sub"])]
for i in range(1, n + 1):
    notes = P.slides[i - 1].notes_slide.notes_text_frame.text if P.slides[i - 1].has_notes_slide else ""
    lines = [l for l in notes.splitlines() if l.strip()]
    m = re.match(r"\[(.*?)\](.*)", lines[0]) if lines else None
    head = f"Slide {i}" + (f"  ·  {m.group(1)}{m.group(2)}" if m else "")
    paras = []
    for l in (lines[1:] if m else lines):
        e = html.escape(l)
        if l.startswith("IF ASKED"):
            paras.append(Paragraph(e, st["ask"]))
        elif l.startswith(("PLAN", "CHECKPOINTS", "DEMO", "TODO")) or l.startswith("~"):
            paras.append(Paragraph(e, st["plan"]))
        else:
            paras.append(Paragraph(e, st["say"]))
    img = Image(f"{thumbs}/s-{i:02d}.png", width=3.0 * inch, height=3.0 * inch * 9 / 16)
    t = Table([[img, paras]], colWidths=[3.15 * inch, 4.35 * inch])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                           ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    story += [KeepTogether([Paragraph(head, st["head"]), t]), Spacer(1, 8)]
SimpleDocTemplate(out, pagesize=letter, leftMargin=0.5 * inch, rightMargin=0.5 * inch,
                  topMargin=0.5 * inch, bottomMargin=0.5 * inch, title="Neural SDEs: Practice Script",
                  author="Nishad Wajge").build(story)
print("written", out)
